from datetime import datetime
from pathlib import Path
import json
import logging
import os
import shutil
from uuid import uuid4
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .activity_record import ActivityRecord
from shared.paths import EXCEL_FILE, USER_DATA_DIR


logger = logging.getLogger(__name__)

BASE_DIR = USER_DATA_DIR
PENDING_FILE = BASE_DIR / "pending_activity_records.jsonl"


def create_excel_file() -> None:
    BASE_DIR.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    work_sheet = workbook.active
    work_sheet.title = "Work Hours"
    work_sheet.append(
        ["Date", "Application", "Start", "End", "Duration (Minutes)"]
    )

    totals_sheet = workbook.create_sheet("Daily Totals")
    totals_sheet.append(
        ["Date", "Total Work Minutes", "Total Work Time"]
    )

    temporary_file = EXCEL_FILE.with_name(
        f"{EXCEL_FILE.stem}_{uuid4().hex}.tmp.xlsx"
    )

    try:
        workbook.save(temporary_file)
        os.replace(temporary_file, EXCEL_FILE)
    finally:
        workbook.close()
        if temporary_file.exists():
            temporary_file.unlink()

    logger.info("Excel file created: %s", EXCEL_FILE)


def backup_corrupted_file() -> None:
    if not EXCEL_FILE.exists():
        return

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_file = EXCEL_FILE.parent / f"work_hours_corrupted_{timestamp}.xlsx"

    shutil.copy2(EXCEL_FILE, backup_file)
    logger.warning("Corrupted Excel file backed up to: %s", backup_file)


def load_or_create_workbook():
    BASE_DIR.mkdir(parents=True, exist_ok=True)

    if not EXCEL_FILE.exists():
        create_excel_file()

    try:
        return load_workbook(EXCEL_FILE)
    except (BadZipFile, InvalidFileException, ValueError):
        logger.exception("Excel file is corrupted: %s", EXCEL_FILE)
        backup_corrupted_file()
        create_excel_file()
        return load_workbook(EXCEL_FILE)


def record_to_dict(record: ActivityRecord) -> dict:
    return {
        "date": record.date,
        "application": record.application,
        "start_time": record.start_time.isoformat(),
        "end_time": record.end_time.isoformat(),
        "duration_seconds": record.duration_seconds,
        "status": record.status,
    }


def append_to_pending(record: ActivityRecord) -> None:
    BASE_DIR.mkdir(parents=True, exist_ok=True)

    with PENDING_FILE.open("a", encoding="utf-8") as file:
        file.write(
            json.dumps(record_to_dict(record), ensure_ascii=False) + "\n"
        )
        file.flush()
        os.fsync(file.fileno())

    logger.warning(
        "Activity record queued because Excel could not be updated: %s",
        record.application,
    )


def append_record_to_sheet(work_sheet, record_data: dict) -> None:
    start_time = datetime.fromisoformat(record_data["start_time"])
    end_time = datetime.fromisoformat(record_data["end_time"])

    work_sheet.append(
        [
            record_data["date"],
            record_data["application"],
            start_time.strftime("%H:%M:%S"),
            end_time.strftime("%H:%M:%S"),
            float(record_data["duration_seconds"]) / 60,
        ]
    )


def update_daily_totals(workbook) -> None:
    if "Work Hours" not in workbook.sheetnames:
        return

    if "Daily Totals" in workbook.sheetnames:
        totals_sheet = workbook["Daily Totals"]
        if totals_sheet.max_row:
            totals_sheet.delete_rows(1, totals_sheet.max_row)
    else:
        totals_sheet = workbook.create_sheet("Daily Totals")

    totals_sheet.append(
        ["Date", "Total Work Minutes", "Total Work Time"]
    )

    totals = {}
    work_sheet = workbook["Work Hours"]

    for row in work_sheet.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 5:
            continue

        date, application, _, _, duration_minutes = row[:5]

        if date is None or duration_minutes is None:
            continue

        if application == "Idle":
            continue

        try:
            duration_minutes = float(duration_minutes)
        except (TypeError, ValueError):
            logger.warning("Skipping invalid duration: %s", duration_minutes)
            continue

        if duration_minutes < 0:
            continue

        totals[date] = totals.get(date, 0.0) + duration_minutes

    for date, total_minutes in sorted(
        totals.items(), key=lambda item: str(item[0])
    ):
        rounded_minutes = round(total_minutes)
        hours, minutes = divmod(rounded_minutes, 60)

        totals_sheet.append(
            [
                date,
                round(total_minutes, 2),
                f"{hours} שעות ו-{minutes} דקות",
            ]
        )


def save_records_to_excel(records: list[dict]) -> None:
    workbook = load_or_create_workbook()
    temporary_file = EXCEL_FILE.with_name(
        f"{EXCEL_FILE.stem}_{uuid4().hex}.tmp.xlsx"
    )

    try:
        work_sheet = workbook["Work Hours"]

        for record_data in records:
            append_record_to_sheet(work_sheet, record_data)

        update_daily_totals(workbook)
        workbook.save(temporary_file)
        workbook.close()
        os.replace(temporary_file, EXCEL_FILE)

    except Exception:
        try:
            workbook.close()
        except Exception:
            pass
        raise

    finally:
        if temporary_file.exists():
            try:
                temporary_file.unlink()
            except OSError:
                logger.warning("Could not remove temporary file: %s", temporary_file)


def read_pending_records() -> list[dict]:
    if not PENDING_FILE.exists():
        return []

    records = []

    with PENDING_FILE.open("r", encoding="utf-8") as file:
        for line_number, line in enumerate(file, start=1):
            if not line.strip():
                continue

            try:
                records.append(json.loads(line))
            except json.JSONDecodeError:
                logger.exception(
                    "Invalid pending record at line %s in %s",
                    line_number,
                    PENDING_FILE,
                )
                raise

    return records


def flush_pending_records() -> None:
    records = read_pending_records()

    if not records:
        return

    save_records_to_excel(records)

    PENDING_FILE.unlink(missing_ok=True)

    logger.info("Saved %s pending activity records", len(records))


def save_activity_record(record: ActivityRecord) -> None:
    record_data = record_to_dict(record)

    try:
        flush_pending_records()
        save_records_to_excel([record_data])

        logger.info(
            "Activity record saved: %s | %.4f minutes | %s",
            record.application,
            record.duration_seconds / 60,
            record.status,
        )

    except (OSError, PermissionError, ValueError, BadZipFile, InvalidFileException):
        logger.exception(
            "Could not save activity record to Excel; queuing it: %s",
            record.application,
        )

        try:
            append_to_pending(record)
        except OSError:
            logger.exception(
                "Could not persist pending activity record: %s",
                record.application,
            )
            raise